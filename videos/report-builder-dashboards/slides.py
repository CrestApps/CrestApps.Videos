"""The slides of the Report Builder dashboards video."""
from pathlib import Path

from PIL import Image, ImageDraw

from brand import BORDER, CODE_BACKGROUND, MARGIN, MUTED, PANEL, TEXT, W, bullets, font, mono, slide_base

DOWNLOADS = Path(__file__).resolve().parent / "shots" / "downloads"

WHY = [
    ("1", "One query, one grouping", "Totals per agent and a trend per day need two groupings"),
    ("2", "Queries, containers, links", "Each visual has its query, a frame, and what a click filters"),
    ("3", "The next question is a click", "Which agent, which queue, which day, and then which calls"),
]

USES = [
    ("1", "A service review", "Scorecards per agent and per queue, the calls one click away"),
    ("2", "A quality check", "A trend by day: click a bad day to list its calls"),
    ("3", "A hand-off to a manager", "Export to Excel: one worksheet for each container"),
]


def _window(image, box, title):
    x0, y0, x1, y1 = box
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(box, radius=14, fill=CODE_BACKGROUND, outline=BORDER, width=2)
    draw.rounded_rectangle((x0, y0, x1, y0 + 52), radius=14, fill=PANEL)
    draw.rectangle((x0, y0 + 30, x1, y0 + 52), fill=PANEL)
    for index, color in enumerate([(237, 106, 94), (245, 191, 79), (98, 197, 84)]):
        draw.ellipse((x0 + 22 + index * 26, y0 + 19, x0 + 36 + index * 26, y0 + 33), fill=color)
    draw.text((x0 + 112, y0 + 26), title, font=font(21, "SemiBold"), fill=MUTED, anchor="lm")
    return draw


def _find(extension):
    files = sorted(DOWNLOADS.glob(f"*{extension}"))
    if not files:
        raise SystemExit(f"No downloaded {extension} file in {DOWNLOADS}")
    return files[0]


def _xlsx_slide(image):
    import openpyxl

    path = _find(".xlsx")
    workbook = openpyxl.load_workbook(path, read_only=True)
    box = (MARGIN, 240, W - MARGIN, 1020)
    draw = _window(image, box, path.name)
    x = box[0] + 30
    y = box[1] + 80
    for index, name in enumerate(workbook.sheetnames):
        width = int(draw.textlength(name, font=font(26, "SemiBold"))) + 48
        draw.rounded_rectangle((x, y, x + width, y + 56), radius=10, fill=(234, 164, 41) if index == 0 else PANEL)
        draw.text((x + 24, y + 28), name, font=font(26, "SemiBold"), fill=(8, 27, 38) if index == 0 else MUTED, anchor="lm")
        x += width + 14
    sheet = workbook[workbook.sheetnames[0]]
    y += 96
    for row in sheet.iter_rows(min_row=1, max_row=9, values_only=True):
        values = [str(value) if value is not None else "" for value in row][:5]
        x = box[0] + 30
        for value in values:
            draw.text((x, y), value[:26], font=font(26), fill=TEXT)
            x += 330
        y += 52
    draw.text((box[0] + 30, box[3] - 50), f"{len(workbook.sheetnames)} worksheets in this workbook", font=font(26, "SemiBold"), fill=MUTED, anchor="lm")


def _pdf_slide(image):
    import fitz

    path = _find(".pdf")
    document = fitz.open(path)
    page = document[0]
    box_height = 1020 - 240
    zoom = box_height / page.rect.height * 2
    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
    page_image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
    page_image = page_image.resize((int(page_image.width / 2), int(page_image.height / 2)))
    x = (W - page_image.width) // 2
    ImageDraw.Draw(image).rounded_rectangle((x - 6, 234, x + page_image.width + 6, 1026), radius=8, fill=BORDER)
    image.paste(page_image, (x, 240))


def slide(step_id, chapter_title, label):
    if step_id.startswith("why-"):
        image = slide_base(chapter_title, label, "A page that answers the next question")
        bullets(image, WHY, active={int(step_id[-1]) - 1})
        return image
    if step_id == "file-1":
        image = slide_base(chapter_title, label, "The Excel workbook")
        _xlsx_slide(image)
        return image
    if step_id == "file-2":
        image = slide_base(chapter_title, label, "The PDF")
        _pdf_slide(image)
        return image
    if step_id.startswith("use-"):
        image = slide_base(chapter_title, label, "Where dashboards fit")
        bullets(image, USES, active={int(step_id[-1]) - 1})
        return image
    raise SystemExit(f"No slide for {step_id}")
